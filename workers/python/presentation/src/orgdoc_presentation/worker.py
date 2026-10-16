from __future__ import annotations

import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

from PIL import Image as PILImage
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

PROTOCOL_VERSION = 1
PPTX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MAX_SLIDES = 40
MAX_BLOCKS = 120
MAX_ASSETS = 100
MAX_TEXT_LENGTH = 4_000
ASSET_ID = re.compile(r"^[A-Za-z][A-Za-z0-9_-]{0,63}$")
HEX_COLOR = re.compile(r"^#?[0-9A-Fa-f]{6}$")
LIBREOFFICE_UNAVAILABLE = (
  "LibreOffice is required for the requested PDF preview. Install LibreOffice or disable "
  "validation.render_pdf."
)
CHART_TYPES = {
  "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
  "line": XL_CHART_TYPE.LINE_MARKERS,
  "pie": XL_CHART_TYPE.PIE,
}
CONTENT_BLOCKS = {"text", "image", "table", "chart"}
LAYOUT_BLOCK_COUNTS = {
  "title": (1, {"title"}),
  "section": (1, {"section"}),
  "content": (2, {"title"}),
  "two_column": (3, {"title"}),
}
FONT_FALLBACK = "Aptos"


class LibreOfficeUnavailable(Exception):
  """The optional PDF preview engine is not installed."""


def _failed(request: dict[str, Any], code: str, message: str) -> dict[str, Any]:
  return {
    "protocol_version": PROTOCOL_VERSION,
    "job_id": request.get("job_id", ""),
    "status": "failed",
    "outputs": [],
    "diagnostics": [{"code": code, "message": message}],
  }


def _color(value: Any, fallback: str) -> RGBColor:
  if isinstance(value, str) and HEX_COLOR.fullmatch(value):
    return RGBColor.from_string(value.removeprefix("#").upper())
  return RGBColor.from_string(fallback)


def _plain_text(value: Any, field: str, maximum: int = MAX_TEXT_LENGTH) -> str:
  if not isinstance(value, str) or not value.strip():
    raise ValueError(f"{field} must be non-empty text")
  if len(value) > maximum:
    raise ValueError(f"{field} exceeds the {maximum} character limit")
  return value


def _validate_document(
  document: dict[str, Any],
) -> tuple[list[dict[str, Any]], dict[str, Path], bool]:
  _plain_text(document.get("title"), "document title", 255)
  slides = document.get("slides")
  if not isinstance(slides, list) or not slides:
    raise ValueError("document must contain at least one slide")
  if len(slides) > MAX_SLIDES:
    raise ValueError(f"document exceeds the limit of {MAX_SLIDES} slides")
  assets = document.get("assets", [])
  if not isinstance(assets, list) or len(assets) > MAX_ASSETS:
    raise ValueError(f"assets must be an array of at most {MAX_ASSETS} items")
  validation = document.get("validation", {})
  if not isinstance(validation, dict):
    raise ValueError("validation must be an object")
  render_pdf = validation.get("render_pdf", False)
  if not isinstance(render_pdf, bool):
    raise ValueError("validation.render_pdf must be true or false")

  root_value = document.get("organization_root")
  root = Path(root_value).resolve() if isinstance(root_value, str) and root_value else None
  asset_paths: dict[str, Path] = {}
  for asset in assets:
    if not isinstance(asset, dict):
      raise ValueError("assets must contain objects with id and path")
    asset_id, relative = asset.get("id"), asset.get("path")
    if not isinstance(asset_id, str) or not ASSET_ID.fullmatch(asset_id) or asset_id in asset_paths:
      raise ValueError(f"invalid or duplicate asset id: {asset_id!r}")
    if root is None:
      raise ValueError("organization_root is required when media assets are declared")
    if not isinstance(relative, str) or not relative or Path(relative).is_absolute():
      raise ValueError(f"asset {asset_id!r} path must be relative to the organization")
    candidate = (root / relative).resolve()
    if not candidate.is_relative_to(root):
      raise ValueError(f"asset {asset_id!r} must be a regular file inside the organization")
    if not candidate.is_file():
      raise ValueError(f"asset {asset_id!r} does not exist inside the organization")
    if candidate.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
      raise ValueError(f"asset {asset_id!r} must use PNG or JPEG format")
    asset_paths[asset_id] = candidate

  block_count = 0
  for index, slide in enumerate(slides, start=1):
    if not isinstance(slide, dict):
      raise ValueError(f"slide {index} must be an object")
    layout, blocks = slide.get("layout"), slide.get("blocks")
    if layout not in LAYOUT_BLOCK_COUNTS:
      raise ValueError(f"slide {index} uses an unsupported layout: {layout!r}")
    if not isinstance(blocks, list):
      raise ValueError(f"slide {index} blocks must be an array")
    expected_count, required_prefix = LAYOUT_BLOCK_COUNTS[layout]
    if len(blocks) != expected_count:
      raise ValueError(f"slide {index} layout {layout!r} requires exactly {expected_count} blocks")
    if any(not isinstance(block, dict) for block in blocks):
      raise ValueError(f"slide {index} blocks must be objects")
    block_types = [block.get("type") for block in blocks]
    for block in blocks:
      if block.get("type") not in {"title", "section", *CONTENT_BLOCKS}:
        raise ValueError(f"slide {index} has unsupported block type: {block.get('type')!r}")
    if block_types[0] not in required_prefix:
      raise ValueError(
        f"slide {index} layout {layout!r} requires a {next(iter(required_prefix))} block first"
      )
    if layout == "content" and block_types[1] not in CONTENT_BLOCKS:
      raise ValueError(f"slide {index} content layout needs one text, image, table, or chart block")
    if layout == "two_column" and any(
      block_type not in CONTENT_BLOCKS for block_type in block_types[1:]
    ):
      raise ValueError(f"slide {index} two_column layout needs two content blocks")
    for block in blocks:
      _validate_block(block, asset_paths, index)
      block_count += 1
  if block_count > MAX_BLOCKS:
    raise ValueError(f"document exceeds the limit of {MAX_BLOCKS} slide blocks")
  return slides, asset_paths, render_pdf


def _validate_block(block: dict[str, Any], assets: dict[str, Path], slide_number: int) -> None:
  block_type = block.get("type")
  label = f"slide {slide_number} {block_type or 'unknown'} block"
  if block_type == "title":
    _plain_text(block.get("text"), f"{label} text", 180)
    if "subtitle" in block:
      _plain_text(block["subtitle"], f"{label} subtitle", 240)
  elif block_type == "section":
    _plain_text(block.get("text"), f"{label} text", 180)
    if "subtitle" in block:
      _plain_text(block["subtitle"], f"{label} subtitle", 300)
  elif block_type == "text":
    heading = block.get("heading")
    if heading is not None:
      _plain_text(heading, f"{label} heading", 120)
    paragraphs = block.get("paragraphs")
    if not isinstance(paragraphs, list) or not paragraphs or len(paragraphs) > 8:
      raise ValueError(f"{label} needs 1-8 paragraphs")
    for paragraph in paragraphs:
      _plain_text(paragraph, f"{label} paragraph", 500)
  elif block_type == "image":
    asset_id = block.get("asset_id")
    if asset_id not in assets:
      raise ValueError(f"{label} must reference an asset declared in document.assets")
    _plain_text(block.get("alt_text"), f"{label} alt_text", 240)
    if "caption" in block:
      _plain_text(block["caption"], f"{label} caption", 180)
  elif block_type == "table":
    columns, rows = block.get("columns"), block.get("rows")
    if not isinstance(columns, list) or not 1 <= len(columns) <= 6:
      raise ValueError(f"{label} must declare 1-6 columns")
    for column in columns:
      _plain_text(column, f"{label} column heading", 80)
    if not isinstance(rows, list) or not 1 <= len(rows) <= 8:
      raise ValueError(f"{label} must contain 1-8 data rows")
    for row in rows:
      if not isinstance(row, list) or len(row) != len(columns):
        raise ValueError(f"{label} rows must match the declared columns")
      for value in row:
        if isinstance(value, bool):
          continue
        if not isinstance(value, (str, int, float)) or (
          isinstance(value, float) and not math.isfinite(value)
        ):
          raise ValueError(f"{label} cells must be text, finite numbers, or booleans")
        if isinstance(value, str) and len(value) > 160:
          raise ValueError(f"{label} cell text exceeds the 160 character limit")
  elif block_type == "chart":
    if block.get("chart_type") not in CHART_TYPES:
      raise ValueError(f"{label} uses an unsupported chart type")
    _plain_text(block.get("title"), f"{label} title", 120)
    categories, series = block.get("categories"), block.get("series")
    if not isinstance(categories, list) or not 2 <= len(categories) <= 12:
      raise ValueError(f"{label} needs 2-12 categories")
    for category in categories:
      _plain_text(category, f"{label} category", 80)
    if not isinstance(series, list) or not 1 <= len(series) <= 4:
      raise ValueError(f"{label} needs 1-4 data series")
    for item in series:
      if not isinstance(item, dict):
        raise ValueError(f"{label} series must be objects")
      _plain_text(item.get("name"), f"{label} series name", 80)
      values = item.get("values")
      if not isinstance(values, list) or len(values) != len(categories):
        raise ValueError(f"{label} series values must match the categories")
      if any(
        isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value)
        for value in values
      ):
        raise ValueError(f"{label} series values must be finite numbers")
  else:
    raise ValueError(f"slide {slide_number} has unsupported block type: {block_type!r}")


def _add_text(
  slide: Any,
  name: str,
  text: str,
  x: float,
  y: float,
  width: float,
  height: float,
  font_name: str,
  font_size: float,
  color: RGBColor,
  bold: bool = False,
  alignment: PP_ALIGN = PP_ALIGN.LEFT,
) -> Any:
  shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
  shape.name = name
  frame = shape.text_frame
  frame.clear()
  frame.word_wrap = True
  frame.vertical_anchor = MSO_ANCHOR.MIDDLE
  frame.margin_left = Inches(0.04)
  frame.margin_right = Inches(0.04)
  frame.margin_top = Inches(0.02)
  frame.margin_bottom = Inches(0.02)
  paragraph = frame.paragraphs[0]
  paragraph.alignment = alignment
  run = paragraph.add_run()
  run.text = text
  run.font.name = font_name or FONT_FALLBACK
  run.font.size = Pt(font_size)
  run.font.bold = bold
  run.font.color.rgb = color
  return shape


def _brand_logo_path(request: dict[str, Any]) -> Path | None:
  payload, context = request["input"], request["input"]["context"]
  root = Path(payload["organization_root"]).resolve()
  candidates = context["brand"].get("assets", [])
  if not isinstance(candidates, list):
    return None
  for relative in candidates:
    if not isinstance(relative, str) or Path(relative).is_absolute():
      continue
    if Path(relative).suffix.lower() not in {".png", ".jpg", ".jpeg"}:
      continue
    candidate = (root / relative).resolve()
    if candidate.is_relative_to(root) and candidate.is_file():
      return candidate
  return None


def _add_logo(slide: Any, path: Path | None) -> None:
  if path is None:
    return
  with PILImage.open(path) as image:
    pixel_width, pixel_height = image.size
  max_width, max_height = 0.9, 0.55
  scale = min(max_width / pixel_width, max_height / pixel_height)
  picture = slide.shapes.add_picture(
    str(path),
    Inches(11.9),
    Inches(0.28),
    width=Inches(pixel_width * scale),
    height=Inches(pixel_height * scale),
  )
  picture.name = "Organization brand logo"
  picture._element.nvPicPr.cNvPr.set("descr", "Organization brand logo")


def _add_slide_title(
  slide: Any, block: dict[str, Any], palette: dict[str, RGBColor], heading_font: str
) -> None:
  _add_text(
    slide,
    "OrgDoc Slide Title",
    block["text"],
    0.72,
    0.42,
    10.8,
    0.68,
    heading_font,
    32,
    palette["primary"],
    bold=True,
  )


def _add_text_block(
  slide: Any,
  block: dict[str, Any],
  x: float,
  y: float,
  width: float,
  height: float,
  palette: dict[str, RGBColor],
  heading_font: str,
  body_font: str,
) -> None:
  top = y
  if block.get("heading"):
    _add_text(
      slide,
      block["heading"],
      block["heading"],
      x,
      top,
      width,
      0.48,
      heading_font,
      20,
      palette["secondary"],
      bold=True,
    )
    top += 0.55
    height -= 0.55
  body = slide.shapes.add_textbox(Inches(x), Inches(top), Inches(width), Inches(height))
  body.name = f"{block['heading']} Content" if block.get("heading") else "OrgDoc Text Content"
  frame = body.text_frame
  frame.clear()
  frame.word_wrap = True
  frame.vertical_anchor = MSO_ANCHOR.TOP
  frame.margin_left = Inches(0.06)
  frame.margin_right = Inches(0.06)
  for index, text in enumerate(block["paragraphs"]):
    paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
    paragraph.text = text
    paragraph.space_after = Pt(12)
    paragraph.font.name = body_font or FONT_FALLBACK
    paragraph.font.size = Pt(18)
    paragraph.font.color.rgb = palette["text"]


def _add_image_block(
  slide: Any,
  block: dict[str, Any],
  asset_paths: dict[str, Path],
  x: float,
  y: float,
  width: float,
  height: float,
  palette: dict[str, RGBColor],
  body_font: str,
) -> None:
  path = asset_paths[block["asset_id"]]
  with PILImage.open(path) as image:
    pixel_width, pixel_height = image.size
  image_box_height = height - (0.55 if block.get("caption") else 0)
  scale = min(width / pixel_width, image_box_height / pixel_height)
  image_width, image_height = pixel_width * scale, pixel_height * scale
  picture = slide.shapes.add_picture(
    str(path),
    Inches(x + (width - image_width) / 2),
    Inches(y + (image_box_height - image_height) / 2),
    width=Inches(image_width),
    height=Inches(image_height),
  )
  picture.name = block["alt_text"]
  picture._element.nvPicPr.cNvPr.set("descr", block["alt_text"])
  if block.get("caption"):
    _add_text(
      slide,
      "OrgDoc Image Caption",
      block["caption"],
      x,
      y + height - 0.5,
      width,
      0.42,
      body_font,
      14,
      palette["secondary"],
      alignment=PP_ALIGN.CENTER,
    )


def _add_table_block(
  slide: Any,
  block: dict[str, Any],
  x: float,
  y: float,
  width: float,
  height: float,
  palette: dict[str, RGBColor],
  body_font: str,
) -> None:
  frame = slide.shapes.add_table(
    len(block["rows"]) + 1,
    len(block["columns"]),
    Inches(x),
    Inches(y),
    Inches(width),
    Inches(height),
  )
  frame.name = "OrgDoc Data Table"
  table = frame.table
  headers = block["columns"]
  values = [headers, *block["rows"]]
  for row_index, row in enumerate(values):
    for column_index, value in enumerate(row):
      cell = table.cell(row_index, column_index)
      cell.text = str(value)
      cell.margin_left = Inches(0.08)
      cell.margin_right = Inches(0.08)
      cell.margin_top = Inches(0.04)
      cell.margin_bottom = Inches(0.04)
      cell.fill.solid()
      cell.fill.fore_color.rgb = palette["primary"] if row_index == 0 else RGBColor(255, 255, 255)
      for paragraph in cell.text_frame.paragraphs:
        paragraph.font.name = body_font or FONT_FALLBACK
        paragraph.font.size = Pt(16 if row_index == 0 else 15)
        paragraph.font.bold = row_index == 0
        paragraph.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else palette["text"]


def _add_chart_block(
  slide: Any,
  block: dict[str, Any],
  x: float,
  y: float,
  width: float,
  height: float,
  palette: dict[str, RGBColor],
  body_font: str,
) -> None:
  chart_data = CategoryChartData()
  chart_data.categories = block["categories"]
  for series in block["series"]:
    chart_data.add_series(series["name"], series["values"])
  shape = slide.shapes.add_chart(
    CHART_TYPES[block["chart_type"]],
    Inches(x),
    Inches(y),
    Inches(width),
    Inches(height),
    chart_data,
  )
  shape.name = "OrgDoc Native Chart"
  chart = shape.chart
  chart.has_title = True
  chart.chart_title.text_frame.text = block["title"]
  title_run = chart.chart_title.text_frame.paragraphs[0].runs[0]
  title_run.font.name = body_font or FONT_FALLBACK
  title_run.font.size = Pt(18)
  title_run.font.bold = True
  title_run.font.color.rgb = palette["primary"]
  chart.has_legend = len(block["series"]) > 1
  if chart.has_legend:
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
  for series_index, series in enumerate(chart.series):
    color = (palette["primary"], palette["secondary"], palette["accent"], palette["text"])[
      series_index % 4
    ]
    if block["chart_type"] == "line":
      series.format.line.color.rgb = color
      series.format.line.width = Pt(2.5)
    elif block["chart_type"] == "pie":
      for point_index, point in enumerate(series.points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = (
          palette["primary"],
          palette["secondary"],
          palette["accent"],
          palette["text"],
        )[point_index % 4]
    else:
      series.format.fill.solid()
      series.format.fill.fore_color.rgb = color


def _render_slides(
  presentation: Presentation,
  slides: list[dict[str, Any]],
  asset_paths: dict[str, Path],
  logo_path: Path | None,
  palette: dict[str, RGBColor],
  heading_font: str,
  body_font: str,
  brand_name: str,
) -> None:
  for slide_spec in slides:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    layout, blocks = slide_spec["layout"], slide_spec["blocks"]
    _add_logo(slide, logo_path)
    if layout == "title":
      block = blocks[0]
      _add_text(
        slide,
        "OrgDoc Slide Title",
        block["text"],
        0.9,
        2.35,
        10.9,
        1.1,
        heading_font,
        42,
        palette["primary"],
        bold=True,
      )
      if block.get("subtitle"):
        _add_text(
          slide,
          "OrgDoc Slide Subtitle",
          block["subtitle"],
          0.95,
          3.55,
          10.5,
          0.7,
          body_font,
          22,
          palette["secondary"],
        )
      _add_text(
        slide, "OrgDoc Brand Name", brand_name, 0.9, 6.78, 10.0, 0.3, body_font, 12, palette["text"]
      )
    elif layout == "section":
      block = blocks[0]
      _add_text(
        slide,
        "OrgDoc Slide Title",
        block["text"],
        0.85,
        2.25,
        11.2,
        0.9,
        heading_font,
        36,
        palette["primary"],
        bold=True,
      )
      if block.get("subtitle"):
        _add_text(
          slide,
          "OrgDoc Section Subtitle",
          block["subtitle"],
          0.9,
          3.28,
          10.7,
          0.65,
          body_font,
          20,
          palette["secondary"],
        )
    else:
      title_block = blocks[0]
      _add_slide_title(slide, title_block, palette, heading_font)
      content = blocks[1:]
      column_width = 5.55 if layout == "two_column" else 11.85
      positions = [(0.78, 1.55), (6.98, 1.55)] if layout == "two_column" else [(0.78, 1.55)]
      for block, (x, y) in zip(content, positions, strict=True):
        width, height = column_width, 5.25
        if block["type"] == "text":
          _add_text_block(slide, block, x, y, width, height, palette, heading_font, body_font)
        elif block["type"] == "image":
          _add_image_block(slide, block, asset_paths, x, y, width, height, palette, body_font)
        elif block["type"] == "table":
          table_height = min(4.3, max(1.2, 0.5 * (len(block["rows"]) + 1)))
          _add_table_block(slide, block, x, y + 0.25, width, table_height, palette, body_font)
        elif block["type"] == "chart":
          _add_chart_block(slide, block, x, y + 0.1, width, 4.75, palette, body_font)


def _render_pdf_preview(presentation_path: Path, staging_dir: Path) -> Path:
  executable = shutil.which("libreoffice") or shutil.which("soffice")
  if executable is None:
    raise LibreOfficeUnavailable(LIBREOFFICE_UNAVAILABLE)
  profile = staging_dir / "lo-profile"
  for directory in (
    profile,
    staging_dir / "lo-home",
    staging_dir / "lo-config",
    staging_dir / "lo-cache",
  ):
    directory.mkdir(parents=True, exist_ok=True)
  try:
    completed = subprocess.run(
      [
        executable,
        "--headless",
        f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(staging_dir),
        str(presentation_path),
      ],
      capture_output=True,
      text=True,
      check=False,
      timeout=120,
      env={
        **os.environ,
        "HOME": str(staging_dir / "lo-home"),
        "XDG_CONFIG_HOME": str(staging_dir / "lo-config"),
        "XDG_CACHE_HOME": str(staging_dir / "lo-cache"),
        "SAL_USE_VCLPLUGIN": "gen",
      },
    )
  except subprocess.TimeoutExpired as exc:
    raise ValueError("LibreOffice presentation preview exceeded the 120 second time limit") from exc
  if completed.returncode != 0:
    detail = (completed.stderr or completed.stdout).strip()
    raise ValueError(f"LibreOffice could not render the presentation: {detail[:500]}")
  preview = staging_dir / "presentation.pdf"
  if not preview.is_file() or preview.stat().st_size == 0:
    detail = (completed.stderr or completed.stdout).strip()
    raise ValueError(f"LibreOffice did not produce a PDF presentation preview: {detail[:500]}")
  preview.rename(staging_dir / "presentation-preview.pdf")
  return staging_dir / "presentation-preview.pdf"


def render(request: dict[str, Any]) -> dict[str, Any]:
  try:
    payload = request["input"]
    document, context = payload["document"], payload["context"]
    if not isinstance(document, dict) or not isinstance(context, dict):
      raise ValueError("input must contain document and context objects")
    slides, asset_paths, render_pdf = _validate_document(
      {**document, "organization_root": payload.get("organization_root")}
    )
    if render_pdf and not (shutil.which("libreoffice") or shutil.which("soffice")):
      raise LibreOfficeUnavailable(LIBREOFFICE_UNAVAILABLE)
    visual = context["brand"].get("visual", {})
    colors = visual.get("colors", {})
    palette = {
      "primary": _color(colors.get("primary"), "173F5F"),
      "secondary": _color(colors.get("secondary"), "20639B"),
      "accent": _color(colors.get("accent"), "F6D55C"),
      "text": _color(colors.get("text"), "17212B"),
    }
    typography = visual.get("typography", {})
    heading_font = typography.get("heading") or FONT_FALLBACK
    body_font = typography.get("body") or FONT_FALLBACK
    for font_name in (heading_font, body_font):
      if not isinstance(font_name, str) or len(font_name) > 80:
        raise ValueError(
          "brand typography values must be font family names of at most 80 characters"
        )
    artifact_dir = Path(request["artifact_dir"])
    artifact_dir.mkdir(parents=True, exist_ok=True)
    pptx_path = artifact_dir / "presentation.pptx"
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = document["title"]
    presentation.core_properties.author = str(context["brand"].get("name") or "OrgDoc")
    logo_path = _brand_logo_path(request)
    with tempfile.TemporaryDirectory(prefix=".orgdoc-presentation-", dir=artifact_dir) as temporary:
      staging_dir = Path(temporary)
      staged_pptx = staging_dir / pptx_path.name
      _render_slides(
        presentation,
        slides,
        asset_paths,
        logo_path,
        palette,
        heading_font,
        body_font,
        str(context["brand"].get("name") or "OrgDoc"),
      )
      presentation.save(staged_pptx)
      staged_preview = _render_pdf_preview(staged_pptx, staging_dir) if render_pdf else None
      staged_pptx.replace(pptx_path)
      if staged_preview is not None:
        staged_preview.replace(artifact_dir / staged_preview.name)
    outputs = [{"path": pptx_path.name, "media_type": PPTX_MEDIA_TYPE}]
    if render_pdf:
      outputs.append({"path": "presentation-preview.pdf", "media_type": "application/pdf"})
    return {
      "protocol_version": PROTOCOL_VERSION,
      "job_id": request.get("job_id", ""),
      "status": "succeeded",
      "outputs": outputs,
      "diagnostics": [],
    }
  except LibreOfficeUnavailable as exc:
    return _failed(request, "libreoffice_unavailable", str(exc))
  except Exception as exc:
    return _failed(request, "presentation_render_error", str(exc))


def main() -> int:
  try:
    request = json.load(sys.stdin)
    result = render(request)
  except Exception as exc:
    result = _failed({}, "invalid_request", str(exc))
  json.dump(result, sys.stdout, ensure_ascii=False, separators=(",", ":"))
  sys.stdout.write("\n")
  return 0


if __name__ == "__main__":
  raise SystemExit(main())
