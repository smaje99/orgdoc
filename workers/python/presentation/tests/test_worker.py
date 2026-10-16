from __future__ import annotations

import base64
import shutil
from pathlib import Path
from unittest.mock import patch

from orgdoc_presentation.worker import render
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

TINY_PNG = base64.b64decode(
  "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+nmWQAAAAASUVORK5CYII="
)


def request_for(tmp_path: Path, brand_id: str = "henko") -> dict:
  colors = {
    "henko": {"primary": "#173F5F", "secondary": "#20639B", "accent": "#F6D55C", "text": "#17212B"},
    "henko-academy": {
      "primary": "#6B3FA0",
      "secondary": "#9B72CF",
      "accent": "#F2C14E",
      "text": "#17212B",
    },
  }
  logo_path = tmp_path / "logo.png"
  logo_path.write_bytes(TINY_PNG)
  return {
    "protocol_version": 1,
    "job_id": "presentation-test",
    "operation_id": "consulting",
    "brand_id": brand_id,
    "document_type": "presentation",
    "artifact_dir": str(tmp_path),
    "input": {
      "organization_root": str(tmp_path),
      "document": {
        "title": "Quarterly results",
        "assets": [{"id": "sample-logo", "path": "logo.png"}],
        "slides": [
          {
            "layout": "title",
            "blocks": [
              {"type": "title", "text": "Quarterly results", "subtitle": "Consulting operation"}
            ],
          },
          {"layout": "section", "blocks": [{"type": "section", "text": "Revenue performance"}]},
          {
            "layout": "two_column",
            "blocks": [
              {"type": "title", "text": "Quarterly revenue"},
              {
                "type": "text",
                "heading": "Summary",
                "paragraphs": ["Protected text stays exactly as supplied."],
              },
              {
                "type": "chart",
                "chart_type": "bar",
                "title": "Revenue by quarter",
                "categories": ["Q1", "Q2"],
                "series": [{"name": "Revenue COP", "values": [1200, 1800]}],
              },
            ],
          },
          {
            "layout": "content",
            "blocks": [
              {"type": "title", "text": "Financial detail"},
              {
                "type": "table",
                "columns": ["Quarter", "Revenue"],
                "rows": [["Q1", 1200], ["Q2", 1800]],
              },
            ],
          },
          {
            "layout": "content",
            "blocks": [
              {"type": "title", "text": "Brand mark"},
              {
                "type": "image",
                "asset_id": "sample-logo",
                "alt_text": "Sample brand mark",
                "caption": "Brand asset",
              },
            ],
          },
        ],
      },
      "context": {
        "organization": {"name": "Henko"},
        "operation": {"name": "Consulting"},
        "brand": {
          "id": brand_id,
          "name": brand_id,
          "version": "1.0.0",
          "assets": ["logo.png"],
          "visual": {
            "colors": colors[brand_id],
            "typography": {"heading": "Aptos Display", "body": "Aptos"},
          },
        },
        "document_type": {
          "id": "presentation",
          "version": "1.0.0",
          "protected_fields": ["slides[].blocks[].text"],
        },
      },
    },
  }


def test_render_preserves_slide_order_text_data_and_editable_objects(tmp_path: Path) -> None:
  result = render(request_for(tmp_path))

  assert result["status"] == "succeeded", result["diagnostics"]
  assert result["outputs"] == [
    {
      "path": "presentation.pptx",
      "media_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }
  ]
  presentation = Presentation(tmp_path / "presentation.pptx")
  assert len(presentation.slides) == 5
  titles = [_shape_by_name(slide, "OrgDoc Slide Title").text for slide in presentation.slides]
  assert titles == [
    "Quarterly results",
    "Revenue performance",
    "Quarterly revenue",
    "Financial detail",
    "Brand mark",
  ]
  summary = _shape_by_name(presentation.slides[2], "Summary Content")
  assert summary.text_frame.text == "Protected text stays exactly as supplied."
  chart = next(shape.chart for shape in presentation.slides[2].shapes if shape.has_chart)
  assert chart.chart_title.text_frame.text == "Revenue by quarter"
  assert chart.series[0].name == "Revenue COP"
  assert list(chart.series[0].values) == [1200, 1800]
  table = next(shape.table for shape in presentation.slides[3].shapes if shape.has_table)
  assert table.cell(0, 0).text == "Quarter"
  assert table.cell(2, 1).text == "1800"
  assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in presentation.slides[4].shapes)
  logo = next(
    shape for shape in presentation.slides[0].shapes if shape.name == "Organization brand logo"
  )
  assert logo._element.nvPicPr.cNvPr.get("descr") == "Organization brand logo"


def test_brand_profile_is_applied_per_deck(tmp_path: Path) -> None:
  henko_dir, academy_dir = tmp_path / "henko", tmp_path / "academy"
  henko_dir.mkdir()
  academy_dir.mkdir()
  assert render(request_for(henko_dir, "henko"))["status"] == "succeeded"
  assert render(request_for(academy_dir, "henko-academy"))["status"] == "succeeded"

  henko = Presentation(henko_dir / "presentation.pptx")
  academy = Presentation(academy_dir / "presentation.pptx")
  henko_title = _shape_by_name(henko.slides[0], "OrgDoc Slide Title")
  academy_title = _shape_by_name(academy.slides[0], "OrgDoc Slide Title")
  assert henko_title.text_frame.paragraphs[0].runs[0].font.color.rgb.__str__() == "173F5F"
  assert academy_title.text_frame.paragraphs[0].runs[0].font.color.rgb.__str__() == "6B3FA0"


def test_image_assets_must_resolve_inside_the_supplied_organization(tmp_path: Path) -> None:
  request = request_for(tmp_path)
  request["input"]["document"]["assets"][0]["path"] = "../logo.png"
  result = render(request)

  assert result["status"] == "failed"
  assert "inside the organization" in result["diagnostics"][0]["message"]
  assert not (tmp_path / "presentation.pptx").exists()


def test_unsupported_blocks_and_layouts_fail_before_creating_a_deck(tmp_path: Path) -> None:
  request = request_for(tmp_path)
  request["input"]["document"]["slides"][2]["blocks"][2]["type"] = "video"
  result = render(request)
  assert result["status"] == "failed"
  assert "unsupported block type" in result["diagnostics"][0]["message"]
  assert not (tmp_path / "presentation.pptx").exists()


def test_requested_pdf_preview_requires_libreoffice(tmp_path: Path) -> None:
  request = request_for(tmp_path)
  request["input"]["document"]["validation"] = {"render_pdf": True}
  with patch("orgdoc_presentation.worker.shutil.which", return_value=None):
    result = render(request)
  assert result["status"] == "failed"
  assert result["diagnostics"][0]["code"] == "libreoffice_unavailable"
  assert "LibreOffice" in result["diagnostics"][0]["message"]
  assert not (tmp_path / "presentation.pptx").exists()


def test_libreoffice_renders_pdf_preview_when_available(tmp_path: Path) -> None:
  if not shutil.which("libreoffice") and not shutil.which("soffice"):
    import pytest

    pytest.skip("LibreOffice is optional")
  request = request_for(tmp_path)
  request["input"]["document"]["validation"] = {"render_pdf": True}
  result = render(request)
  assert result["status"] == "succeeded", result["diagnostics"]
  assert result["outputs"][-1] == {
    "path": "presentation-preview.pdf",
    "media_type": "application/pdf",
  }
  assert (tmp_path / "presentation-preview.pdf").read_bytes().startswith(b"%PDF-")


def _shape_by_name(slide, name: str):
  for shape in slide.shapes:
    if shape.name == name:
      return shape
  raise AssertionError(f"shape {name!r} is missing")
