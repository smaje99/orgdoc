from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import tempfile
from copy import copy
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

PROTOCOL_VERSION = 1
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_CELLS = 50_000
MAX_SHEETS = 20
MAX_CHARTS = 100
MAX_FORMULA_LENGTH = 2_000
FORMULA_FUNCTION_DENYLIST = re.compile(
  r"(?i)\b(WEBSERVICE|HYPERLINK|RTD|FILTERXML|IMAGE|CALL|REGISTER\.ID|EXEC)\s*\("
)
FORMULA_EXTERNAL_REFERENCE = re.compile(r"\[[^\]]+\]")
TABLE_NAME = re.compile(r"^[A-Za-z_][A-Za-z0-9_.]*$")
CELL_ADDRESS = re.compile(r"^[A-Z]{1,3}[1-9]\d{0,6}$")
DEFAULT_FORMATS = {
  "currency": '"$"#,##0.00;[Red]("$"#,##0.00)',
  "percentage": "0.00%",
  "date": "yyyy-mm-dd",
}
TABLE_STYLES = {"TableStyleMedium2", "TableStyleMedium4", "TableStyleMedium9", "TableStyleMedium21"}
LIBREOFFICE_UNAVAILABLE = (
  "LibreOffice is required for the requested recalculation or PDF preview. Install LibreOffice, "
  "or choose formula_policy=spreadsheet_app and disable validation.render_pdf."
)


def _failed(request: dict[str, Any], code: str, message: str) -> dict[str, Any]:
  return {
    "protocol_version": PROTOCOL_VERSION,
    "job_id": request.get("job_id", ""),
    "status": "failed",
    "outputs": [],
    "diagnostics": [{"code": code, "message": message}],
  }


def _brand_color(colors: dict[str, Any], name: str, fallback: str) -> str:
  value = colors.get(name)
  if isinstance(value, str) and re.fullmatch(r"#?[0-9A-Fa-f]{6}", value):
    return value.removeprefix("#").upper()
  return fallback


def _validate_formula(formula: Any) -> str:
  if (
    not isinstance(formula, str) or not formula.startswith("=") or len(formula) > MAX_FORMULA_LENGTH
  ):
    raise ValueError(
      "formula cells need a formula beginning with '=' and no more than 2000 characters"
    )
  if FORMULA_EXTERNAL_REFERENCE.search(formula):
    raise ValueError("external workbook references are not supported in formulas")
  if FORMULA_FUNCTION_DENYLIST.search(formula):
    raise ValueError(
      "unsupported function in formula; external-content and macro functions are blocked"
    )
  return formula


def _cell_value(cell: dict[str, Any]) -> Any:
  cell_type = cell.get("type")
  if cell_type == "formula":
    return _validate_formula(cell.get("formula"))
  if "value" not in cell:
    raise ValueError(f"{cell_type} cells require a value")
  value = cell["value"]
  if cell_type in {"number", "currency", "percentage"}:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
      raise ValueError(f"{cell_type} cell values must be JSON numbers")
  elif cell_type == "boolean":
    if not isinstance(value, bool):
      raise ValueError("boolean cell values must be true or false")
  elif cell_type == "string":
    if not isinstance(value, str):
      raise ValueError("string cell values must be text")
  elif cell_type == "date":
    if not isinstance(value, str):
      raise ValueError("date cell values must use an ISO date string")
    try:
      return date.fromisoformat(value)
    except ValueError as exc:
      raise ValueError("date cell values must use YYYY-MM-DD") from exc
  else:
    raise ValueError(f"unsupported cell type: {cell_type!r}")
  return value


def _validate_document(document: dict[str, Any]) -> tuple[list[dict[str, Any]], str, bool]:
  sheets = document.get("sheets")
  if not isinstance(sheets, list) or not sheets:
    raise ValueError("document must contain at least one sheet")
  if len(sheets) > MAX_SHEETS:
    raise ValueError(f"document exceeds the limit of {MAX_SHEETS} sheets")
  formula_policy = document.get("formula_policy")
  if formula_policy not in {"spreadsheet_app", "libreoffice"}:
    raise ValueError("formula_policy must be spreadsheet_app or libreoffice")
  validation = document.get("validation", {})
  if not isinstance(validation, dict):
    raise ValueError("validation must be an object")
  render_pdf = validation.get("render_pdf", False)
  if not isinstance(render_pdf, bool):
    raise ValueError("validation.render_pdf must be true or false")

  names: set[str] = set()
  tables: set[str] = set()
  cell_count = 0
  chart_count = 0
  for sheet in sheets:
    if not isinstance(sheet, dict):
      raise ValueError("each sheet must be an object")
    name, rows = sheet.get("name"), sheet.get("rows")
    if (
      not isinstance(name, str)
      or not name
      or len(name) > 31
      or any(char in name for char in "[]:*?/\\")
    ):
      raise ValueError("sheet names must be 1-31 characters and cannot contain []:*?/\\")
    if name.casefold() in names:
      raise ValueError(f"duplicate sheet name: {name}")
    names.add(name.casefold())
    if not isinstance(rows, list) or not rows:
      raise ValueError(f"sheet {name!r} must contain at least one row")
    for row in rows:
      if not isinstance(row, list) or not row:
        raise ValueError(f"sheet {name!r} rows must be non-empty arrays")
      cell_count += len(row)
      for cell in row:
        if not isinstance(cell, dict):
          raise ValueError(f"sheet {name!r} cells must be typed objects")
        _cell_value(cell)
        number_format = cell.get("number_format")
        if number_format is not None and (
          not isinstance(number_format, str)
          or len(number_format) > 128
          or any(ord(char) < 32 for char in number_format)
        ):
          raise ValueError("number_format must be a printable string of at most 128 characters")
    for table in sheet.get("tables", []):
      if not isinstance(table, dict):
        raise ValueError("tables must be objects")
      table_name, ref = table.get("name"), table.get("ref")
      if (
        not isinstance(table_name, str)
        or not TABLE_NAME.fullmatch(table_name)
        or table_name.casefold() in tables
      ):
        raise ValueError(f"invalid or duplicate table name: {table_name!r}")
      if not isinstance(ref, str):
        raise ValueError(f"table {table_name!r} needs a cell range")
      if table.get("style", "TableStyleMedium2") not in TABLE_STYLES:
        raise ValueError(f"table {table_name!r} uses an unsupported table style")
      try:
        min_col, min_row, max_col, max_row = range_boundaries(ref)
      except (TypeError, ValueError) as exc:
        raise ValueError(f"table {table_name!r} has an invalid range") from exc
      if None in (min_col, min_row, max_col, max_row) or min_row == max_row or min_col == max_col:
        raise ValueError(
          f"table {table_name!r} range needs headers, data rows, and at least two columns"
        )
      if max_row > len(rows) or max_col > max((len(row) for row in rows), default=0):
        raise ValueError(f"table {table_name!r} range extends beyond populated cells")
      header_values = [rows[min_row - 1][column - 1] for column in range(min_col, max_col + 1)]
      if any(item.get("type") != "string" or not item.get("value") for item in header_values):
        raise ValueError(f"table {table_name!r} needs non-empty text headers in its first row")
      if len({item["value"] for item in header_values}) != len(header_values):
        raise ValueError(f"table {table_name!r} column headers must be unique")
      tables.add(table_name.casefold())
    charts = sheet.get("charts", [])
    if not isinstance(charts, list):
      raise ValueError(f"sheet {name!r} charts must be an array")
    chart_count += len(charts)
  if cell_count > MAX_CELLS:
    raise ValueError(f"document exceeds the limit of {MAX_CELLS} cells")
  if chart_count > MAX_CHARTS:
    raise ValueError(f"document exceeds the limit of {MAX_CHARTS} charts")
  return sheets, formula_policy, render_pdf


def _set_cell_style(
  target: Any, source: dict[str, Any], colors: dict[str, str], row_index: int
) -> None:
  cell_type = source.get("type")
  target.number_format = source.get("number_format", DEFAULT_FORMATS.get(cell_type, "General"))
  protection = copy(target.protection)
  protection.locked = source.get("protected", True)
  target.protection = protection
  if row_index == 1:
    target.fill = PatternFill("solid", fgColor=colors["primary"])
    target.font = Font(name="Aptos", bold=True, color="FFFFFF")
    target.alignment = Alignment(vertical="center", wrap_text=True)
  elif cell_type == "formula":
    target.font = Font(name="Aptos", color=colors["secondary"], bold=True)
  elif cell_type in {"currency", "number", "percentage"}:
    target.alignment = Alignment(horizontal="right")


def _add_chart(worksheet: Any, chart: dict[str, Any], colors: dict[str, str]) -> None:
  chart_type = chart.get("type")
  constructors = {"bar": BarChart, "line": LineChart, "pie": PieChart}
  if chart_type not in constructors:
    raise ValueError(
      f"unsupported chart type: {chart_type!r}; supported types are bar, line, and pie"
    )
  data_ref, categories_ref = chart.get("data_ref"), chart.get("categories_ref")
  anchor = chart.get("anchor")
  if not all(isinstance(item, str) for item in (data_ref, categories_ref, anchor)):
    raise ValueError("charts require data_ref, categories_ref, and anchor")
  if not CELL_ADDRESS.fullmatch(anchor):
    raise ValueError("chart anchor must be a cell address such as D2")
  try:
    data_bounds = range_boundaries(data_ref)
    category_bounds = range_boundaries(categories_ref)
  except (TypeError, ValueError) as exc:
    raise ValueError("chart references must be simple cell ranges on the current sheet") from exc
  if None in data_bounds or None in category_bounds:
    raise ValueError("chart references must be simple cell ranges on the current sheet")
  data_min_col, data_min_row, data_max_col, data_max_row = data_bounds
  cat_min_col, cat_min_row, cat_max_col, cat_max_row = category_bounds
  if data_min_col != data_max_col or cat_min_col != cat_max_col or data_max_row < data_min_row:
    raise ValueError("charts support one numeric data column and one category column")
  if (
    data_max_row > worksheet.max_row
    or cat_max_row > worksheet.max_row
    or max(data_max_col, cat_max_col) > worksheet.max_column
  ):
    raise ValueError("chart references extend beyond populated cells")
  if data_max_row - data_min_row != cat_max_row - cat_min_row:
    raise ValueError("chart data and category references must have the same number of rows")
  chart_object = constructors[chart_type]()
  chart_object.title = str(chart.get("title") or "")
  chart_object.add_data(
    Reference(worksheet, min_col=data_min_col, min_row=data_min_row, max_row=data_max_row),
    titles_from_data=False,
  )
  chart_object.set_categories(
    Reference(worksheet, min_col=cat_min_col, min_row=cat_min_row, max_row=cat_max_row)
  )
  if chart_type in {"bar", "line"}:
    if chart_type == "bar":
      chart_object.type = "col"
    for series in chart_object.series:
      series.graphicalProperties.line.solidFill = colors["primary"]
      series.graphicalProperties.solidFill = colors["primary"]
  else:
    chart_object.dataLabels = DataLabelList()
    chart_object.dataLabels.showPercent = True
    chart_object.dataLabels.showLeaderLines = True
    for index, color in enumerate((colors["primary"], colors["secondary"], colors["accent"])):
      from openpyxl.chart.shapes import GraphicalProperties

      chart_object.series[0].data_points.append(
        DataPoint(idx=index, spPr=GraphicalProperties(solidFill=color))
      )
  worksheet.add_chart(chart_object, anchor)


def _add_brand_image(request: dict[str, Any], worksheet: Any) -> None:
  payload = request["input"]
  context = payload["context"]
  root = Path(payload["organization_root"]).resolve()
  for asset in context["brand"].get("assets", []):
    if not isinstance(asset, str) or Path(asset).is_absolute():
      continue
    if Path(asset).suffix.lower() not in {".png", ".jpg", ".jpeg"}:
      continue
    candidate = (root / asset).resolve()
    if candidate.is_relative_to(root) and candidate.is_file():
      image = Image(str(candidate))
      scale = min(1, 220 / image.width, 64 / image.height)
      image.width *= scale
      image.height *= scale
      worksheet.add_image(image, worksheet.cell(1, worksheet.max_column + 2).coordinate)
      return


def _libreoffice_command(
  executable: str, source: Path, target_dir: Path, output_format: str
) -> None:
  target_dir.mkdir(parents=True, exist_ok=True)
  profile = target_dir / "lo-profile"
  profile_uri = profile.resolve().as_uri()
  try:
    completed = subprocess.run(
      [
        executable,
        "--headless",
        f"-env:UserInstallation={profile_uri}",
        "--convert-to",
        output_format,
        "--outdir",
        str(target_dir),
        str(source),
      ],
      capture_output=True,
      text=True,
      check=False,
      timeout=120,
    )
  except subprocess.TimeoutExpired as exc:
    raise ValueError("LibreOffice validation exceeded the 120 second time limit") from exc
  if completed.returncode != 0:
    detail = (completed.stderr or completed.stdout).strip()
    raise ValueError(f"LibreOffice could not convert the workbook: {detail[:500]}")


def _validate_with_libreoffice(
  workbook_path: Path, artifact_dir: Path, recalculate: bool, render_pdf: bool
) -> list[dict[str, str]]:
  executable = shutil.which("libreoffice") or shutil.which("soffice")
  if executable is None:
    raise FileNotFoundError(LIBREOFFICE_UNAVAILABLE)
  outputs: list[dict[str, str]] = []
  if recalculate:
    with tempfile.TemporaryDirectory(prefix="orgdoc-recalc-") as temporary:
      recalculated_dir = Path(temporary)
      _libreoffice_command(executable, workbook_path, recalculated_dir, "xlsx")
      recalculated = recalculated_dir / workbook_path.name
      if not recalculated.is_file() or recalculated.stat().st_size == 0:
        raise ValueError("LibreOffice did not produce a recalculated workbook")
      validated = load_workbook(recalculated, data_only=False, read_only=False)
      validated.close()
      recalculated.replace(workbook_path)
  if render_pdf:
    _libreoffice_command(executable, workbook_path, artifact_dir, "pdf")
    converted_pdf = artifact_dir / f"{workbook_path.stem}.pdf"
    preview = artifact_dir / "spreadsheet-preview.pdf"
    if not converted_pdf.is_file() or converted_pdf.stat().st_size == 0:
      raise ValueError("LibreOffice did not produce a PDF preview")
    converted_pdf.replace(preview)
    outputs.append({"path": preview.name, "media_type": "application/pdf"})
  return outputs


def render(request: dict[str, Any]) -> dict[str, Any]:
  try:
    payload = request["input"]
    document = payload["document"]
    context = payload["context"]
    if not isinstance(document, dict) or not isinstance(context, dict):
      raise ValueError("input must contain document and context objects")
    sheets, formula_policy, render_pdf = _validate_document(document)
    if (formula_policy == "libreoffice" or render_pdf) and not (
      shutil.which("libreoffice") or shutil.which("soffice")
    ):
      raise FileNotFoundError(LIBREOFFICE_UNAVAILABLE)
    brand = context["brand"]
    visual = brand.get("visual", {})
    theme = visual.get("colors", {})
    colors = {
      "primary": _brand_color(theme, "primary", "173F5F"),
      "secondary": _brand_color(theme, "secondary", "20639B"),
      "accent": _brand_color(theme, "accent", "F6D55C"),
      "text": _brand_color(theme, "text", "17212B"),
    }
    artifact_dir = Path(request["artifact_dir"])
    artifact_dir.mkdir(parents=True, exist_ok=True)
    workbook_path = artifact_dir / "spreadsheet.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = str(document.get("title") or "OrgDoc Spreadsheet")
    workbook.properties.creator = str(brand.get("name") or "OrgDoc")
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    for sheet_spec in sheets:
      worksheet = workbook.create_sheet(sheet_spec["name"])
      worksheet.sheet_view.showGridLines = False
      worksheet.sheet_properties.tabColor = colors["accent"]
      worksheet.freeze_panes = "A2"
      worksheet.sheet_view.zoomScale = 90
      for row_index, row in enumerate(sheet_spec["rows"], start=1):
        for column_index, cell_spec in enumerate(row, start=1):
          target = worksheet.cell(row=row_index, column=column_index, value=_cell_value(cell_spec))
          _set_cell_style(target, cell_spec, colors, row_index)
      if sheet_spec.get("protect", False):
        worksheet.protection.sheet = True
        worksheet.protection.selectLockedCells = False
        worksheet.protection.selectUnlockedCells = True
      for table_spec in sheet_spec.get("tables", []):
        table = Table(displayName=table_spec["name"], ref=table_spec["ref"])
        table.tableStyleInfo = TableStyleInfo(
          name=table_spec.get("style", "TableStyleMedium2"),
          showFirstColumn=False,
          showLastColumn=False,
          showRowStripes=True,
          showColumnStripes=False,
        )
        worksheet.add_table(table)
      for chart_spec in sheet_spec.get("charts", []):
        _add_chart(worksheet, chart_spec, colors)
      worksheet.auto_filter.ref = f"A1:{worksheet.cell(1, worksheet.max_column).coordinate}"
      worksheet.row_dimensions[1].height = 24
      for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter
        width = min(
          max(max((len(str(cell.value or "")) for cell in column_cells), default=8) + 2, 10), 42
        )
        worksheet.column_dimensions[column_letter].width = width
    _add_brand_image(request, workbook[workbook.sheetnames[0]])
    with tempfile.TemporaryDirectory(prefix=".orgdoc-spreadsheet-", dir=artifact_dir) as temporary:
      staged_dir = Path(temporary)
      staged_workbook = staged_dir / workbook_path.name
      workbook.save(staged_workbook)
      validation_outputs = (
        _validate_with_libreoffice(
          staged_workbook,
          staged_dir,
          recalculate=formula_policy == "libreoffice",
          render_pdf=render_pdf,
        )
        if formula_policy == "libreoffice" or render_pdf
        else []
      )
      staged_workbook.replace(workbook_path)
      for output in validation_outputs:
        (staged_dir / output["path"]).replace(artifact_dir / output["path"])
    return {
      "protocol_version": PROTOCOL_VERSION,
      "job_id": request.get("job_id", ""),
      "status": "succeeded",
      "outputs": [{"path": workbook_path.name, "media_type": XLSX_MEDIA_TYPE}, *validation_outputs],
      "diagnostics": [],
    }
  except FileNotFoundError as exc:
    return _failed(request, "libreoffice_unavailable", str(exc))
  except Exception as exc:
    return _failed(request, "spreadsheet_render_error", str(exc))


def main() -> int:
  try:
    request = json.load(sys.stdin)
    result = render(request)
  except Exception as exc:
    result = _failed({}, "invalid_request", str(exc))
  json.dump(result, sys.stdout, ensure_ascii=False, separators=(",", ":"))
  sys.stdout.write("\n")
  return 0


if __name__ == "__main__":
  raise SystemExit(main())
