from __future__ import annotations

import shutil
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from orgdoc_spreadsheet.worker import render


def request_for(tmp_path: Path, brand_id: str = "henko") -> dict:
  colors = {
    "henko": {"primary": "#173F5F", "secondary": "#20639B", "accent": "#F6D55C", "text": "#17212B"},
    "henko-academy": {
      "primary": "#6B3FA0",
      "secondary": "#9B72CF",
      "accent": "#F2C14E",
      "text": "#17212B",
    },
  }
  return {
    "protocol_version": 1,
    "job_id": "spreadsheet-test",
    "operation_id": "consulting",
    "brand_id": brand_id,
    "document_type": "spreadsheet",
    "artifact_dir": str(tmp_path),
    "input": {
      "organization_root": str(tmp_path),
      "document": {
        "title": "Quarterly Results",
        "formula_policy": "spreadsheet_app",
        "sheets": [
          {
            "name": "Summary",
            "protect": True,
            "rows": [
              [
                {"type": "string", "value": "Quarter"},
                {"type": "string", "value": "Revenue"},
              ],
              [
                {"type": "string", "value": "Q1"},
                {"type": "currency", "value": 1200, "protected": True},
              ],
              [
                {"type": "string", "value": "Q2"},
                {"type": "currency", "value": 1800, "protected": True},
              ],
              [
                {"type": "string", "value": "Total"},
                {"type": "formula", "formula": "=SUM(B2:B3)", "number_format": '"$"#,##0.00'},
              ],
            ],
            "tables": [{"name": "RevenueData", "ref": "A1:B3", "style": "TableStyleMedium2"}],
            "charts": [
              {
                "type": "bar",
                "title": "Quarterly revenue",
                "data_ref": "B2:B3",
                "categories_ref": "A2:A3",
                "anchor": "D2",
              }
            ],
          }
        ],
      },
      "context": {
        "organization": {"name": "Henko"},
        "operation": {"name": "Consulting"},
        "brand": {
          "id": brand_id,
          "name": brand_id,
          "version": "1.0.0",
          "visual": {"colors": colors[brand_id]},
          "assets": [],
        },
        "document_type": {"id": "spreadsheet", "version": "1.0.0", "protected_fields": ["revenue"]},
      },
    },
  }


def test_render_writes_branded_workbook_with_formulas_tables_charts_and_protection(
  tmp_path: Path,
) -> None:
  result = render(request_for(tmp_path))

  assert result["status"] == "succeeded"
  assert result["outputs"] == [
    {
      "path": "spreadsheet.xlsx",
      "media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
  ]
  workbook = load_workbook(tmp_path / "spreadsheet.xlsx", data_only=False)
  sheet = workbook["Summary"]
  assert workbook.properties.title == "Quarterly Results"
  assert sheet["A1"].fill.fgColor.rgb.endswith("173F5F")
  assert sheet["B2"].value == 1200
  assert sheet["B2"].protection.locked is True
  assert sheet["B4"].value == "=SUM(B2:B3)"
  assert sheet["B4"].number_format == '"$"#,##0.00'
  assert sheet.protection.sheet is True
  assert list(sheet.tables) == ["RevenueData"]
  assert sheet.tables["RevenueData"].ref == "A1:B3"
  assert sheet.tables["RevenueData"].tableStyleInfo.name == "TableStyleMedium2"
  assert len(sheet._charts) == 1
  chart = sheet._charts[0]
  assert chart.type == "col"
  assert chart.series[0].val.numRef.f.endswith("$B$2:$B$3")
  category_reference = chart.series[0].cat.strRef or chart.series[0].cat.numRef
  assert category_reference.f.endswith("$A$2:$A$3")
  assert workbook.calculation.fullCalcOnLoad is True
  assert workbook.calculation.forceFullCalc is True
  workbook.close()
  workbook_values = load_workbook(tmp_path / "spreadsheet.xlsx", data_only=True)
  assert workbook_values["Summary"]["B4"].value is None
  workbook_values.close()


def test_brand_profile_is_applied_per_render_without_shared_style_state(tmp_path: Path) -> None:
  first_dir = tmp_path / "henko"
  second_dir = tmp_path / "academy"
  first_dir.mkdir()
  second_dir.mkdir()

  assert render(request_for(first_dir, "henko"))["status"] == "succeeded"
  assert render(request_for(second_dir, "henko-academy"))["status"] == "succeeded"
  first = load_workbook(first_dir / "spreadsheet.xlsx")
  second = load_workbook(second_dir / "spreadsheet.xlsx")
  assert first["Summary"]["A1"].fill.fgColor.rgb.endswith("173F5F")
  assert second["Summary"]["A1"].fill.fgColor.rgb.endswith("6B3FA0")


def test_line_and_pie_chart_types_are_supported(tmp_path: Path) -> None:
  for chart_type in ("line", "pie"):
    artifact_dir = tmp_path / chart_type
    artifact_dir.mkdir()
    request = request_for(artifact_dir)
    request["input"]["document"]["sheets"][0]["charts"][0]["type"] = chart_type
    result = render(request)
    assert result["status"] == "succeeded", result["diagnostics"]
    workbook = load_workbook(artifact_dir / "spreadsheet.xlsx")
    assert len(workbook["Summary"]._charts) == 1


def test_recalculation_request_fails_clearly_when_libreoffice_is_missing(tmp_path: Path) -> None:
  request = request_for(tmp_path)
  request["input"]["document"]["formula_policy"] = "libreoffice"
  with patch("orgdoc_spreadsheet.worker.shutil.which", return_value=None):
    result = render(request)

  assert result["status"] == "failed"
  assert result["diagnostics"][0]["code"] == "libreoffice_unavailable"
  assert "LibreOffice" in result["diagnostics"][0]["message"]
  assert not (tmp_path / "spreadsheet.xlsx").exists()


def test_libreoffice_recalculates_formulas_and_renders_pdf_when_available(tmp_path: Path) -> None:
  if not shutil.which("libreoffice") and not shutil.which("soffice"):
    import pytest

    pytest.skip("LibreOffice is optional")
  request = request_for(tmp_path)
  request["input"]["document"]["formula_policy"] = "libreoffice"
  request["input"]["document"]["validation"] = {"render_pdf": True}
  result = render(request)

  assert result["status"] == "succeeded", result["diagnostics"]
  assert result["outputs"] == [
    {
      "path": "spreadsheet.xlsx",
      "media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    },
    {"path": "spreadsheet-preview.pdf", "media_type": "application/pdf"},
  ]
  workbook = load_workbook(tmp_path / "spreadsheet.xlsx", data_only=True)
  assert workbook["Summary"]["B4"].value == 3000
  assert (tmp_path / "spreadsheet-preview.pdf").read_bytes().startswith(b"%PDF-")


def test_unsupported_chart_and_formula_are_rejected_before_workbook_is_written(
  tmp_path: Path,
) -> None:
  request = request_for(tmp_path)
  request["input"]["document"]["sheets"][0]["charts"][0]["type"] = "surface"
  result = render(request)
  assert result["status"] == "failed"
  assert "chart type" in result["diagnostics"][0]["message"]
  assert not (tmp_path / "spreadsheet.xlsx").exists()

  request = request_for(tmp_path)
  request["input"]["document"]["sheets"][0]["rows"][3][1]["formula"] = (
    '=WEBSERVICE("https://example.invalid")'
  )
  result = render(request)
  assert result["status"] == "failed"
  assert "unsupported function" in result["diagnostics"][0]["message"]
  assert not (tmp_path / "spreadsheet.xlsx").exists()
